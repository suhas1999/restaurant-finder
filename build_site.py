"""Build a phone-friendly static site in docs/ from results/<Town>_*.xlsx (served by GitHub Pages).

Usage: python3 build_site.py   ->  docs/index.html + docs/<town>/index.html for every town in results/
"""
import html
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
RESULTS, DOCS = HERE / "results", HERE / "docs"

CSS = """
:root{--bg:#f6f4ef;--card:#fff;--text:#1d1d1f;--muted:#6b6b70;--accent:#c2410c;--star:#f59e0b;--line:#e7e3da}
@media (prefers-color-scheme:dark){:root{--bg:#121212;--card:#1e1e1e;--text:#f2f2f2;--muted:#a0a0a5;--accent:#fb923c;--line:#2c2c2c}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font:17px/1.4 -apple-system,system-ui,"Segoe UI",Roboto,sans-serif}
header{padding:20px 16px 8px}h1{margin:0;font-size:26px}header p{margin:4px 0 0;color:var(--muted);font-size:15px}
.tabs{position:sticky;top:0;z-index:2;display:flex;gap:8px;padding:10px 16px;background:var(--bg);border-bottom:1px solid var(--line)}
.tabs button{flex:1;padding:12px;border:0;border-radius:12px;font-size:16px;font-weight:600;background:var(--card);color:var(--text);box-shadow:0 1px 2px rgba(0,0,0,.08)}
.tabs button.on{background:var(--accent);color:#fff}
.note{margin:12px 16px 0;color:var(--muted);font-size:14px}
ol{list-style:none;margin:0;padding:12px 16px 40px}
li{background:var(--card);border-radius:14px;padding:14px;margin-bottom:12px;box-shadow:0 1px 3px rgba(0,0,0,.08);display:flex;gap:12px}
.rank{flex:none;width:34px;height:34px;border-radius:50%;background:var(--accent);color:#fff;font-weight:700;display:grid;place-items:center;font-size:15px}
.body{min-width:0;flex:1}.name{font-weight:700;font-size:18px;overflow-wrap:anywhere}
.meta{color:var(--muted);font-size:14px;margin-top:2px}.stars{color:var(--star);font-weight:700}
.addr{color:var(--muted);font-size:14px;margin-top:6px}
.btns{display:flex;gap:8px;margin-top:10px}.btns a{flex:1;text-align:center;padding:10px;border-radius:10px;text-decoration:none;font-weight:600;font-size:15px;border:1px solid var(--line);color:var(--text)}
.btns a.none{color:var(--muted);font-weight:400}.btns a{white-space:nowrap}.btns a.map{background:var(--accent);color:#fff;border-color:var(--accent)}
.towns a{display:block;margin:12px 16px;padding:18px;background:var(--card);border-radius:14px;text-decoration:none;color:var(--text);font-size:20px;font-weight:700;box-shadow:0 1px 3px rgba(0,0,0,.08)}
.hidden{display:none}
"""

HEAD = """<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title><style>{css}</style></head><body>"""


def read(path):
    rows = list(load_workbook(path).active.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def card(r):
    e = lambda v: html.escape(str(v or ""))
    rating = f'<span class="stars">★ {r["Rating"]}</span> · {int(r["Reviews"] or 0):,} reviews' if r["Rating"] else "No rating"
    km = f' · {r["Km from Centre"]} km' if r.get("Km from Centre") is not None else ""
    phone = str(r["Phone"] or "").strip()
    call = f'<a href="tel:{e(phone.replace(" ", ""))}">📞 {e(phone)}</a>' if phone else '<a class="none">No phone listed</a>'
    return (f'<li><div class="rank">{r["Rank"]}</div><div class="body"><div class="name">{e(r["Restaurant Name"])}</div>'
            f'<div class="meta">{rating}</div><div class="meta">{e(r["Category"])}{km}</div>'
            f'<div class="addr">{e(r["Address"])}</div>'
            f'<div class="btns"><a class="map" href="{e(r["Google Maps Link"])}">📍 Open in Maps</a>{call}</div></div></li>')


def town_page(town, top, pop):
    return (HEAD.format(title=f"{town} Restaurants", css=CSS)
            + f'<header><h1>🍽️ {html.escape(town)}</h1><p>Best restaurants from Google Maps reviews</p></header>'
            + '<div class="tabs"><button class="on" data-t="top">⭐ Top Rated</button><button data-t="pop">🔥 Most Popular</button></div>'
            + '<p class="note" id="n-top">Best ratings, with more weight to places that have many reviews.</p>'
            + '<p class="note hidden" id="n-pop">Most reviewed places rated 4.0 or above.</p>'
            + '<ol id="top">' + "".join(card(r) for r in top) + "</ol>"
            + '<ol id="pop" class="hidden">' + "".join(card(r) for r in pop) + "</ol>"
            + """<script>document.querySelectorAll('.tabs button').forEach(b=>b.onclick=()=>{
document.querySelectorAll('.tabs button').forEach(x=>x.classList.toggle('on',x===b));
['top','pop'].forEach(t=>{document.getElementById(t).classList.toggle('hidden',t!==b.dataset.t);
document.getElementById('n-'+t).classList.toggle('hidden',t!==b.dataset.t)});scrollTo(0,0)})</script></body></html>""")


def main():
    towns = sorted(p.name[:-len("_top_rated.xlsx")] for p in RESULTS.glob("*_top_rated.xlsx") if not p.name.startswith("~$"))
    for t in towns:
        top, pop = read(RESULTS / f"{t}_top_rated.xlsx"), read(RESULTS / f"{t}_most_popular.xlsx")
        out = DOCS / t.lower()
        out.mkdir(parents=True, exist_ok=True)
        (out / "index.html").write_text(town_page(t.replace("_", " "), top, pop))
        print(f"docs/{t.lower()}/index.html: {len(top)} top rated, {len(pop)} most popular")
    links = "".join(f'<a href="{t.lower()}/">{html.escape(t.replace("_", " "))} →</a>' for t in towns)
    (DOCS / "index.html").write_text(HEAD.format(title="Restaurant Finder", css=CSS)
                                     + '<header><h1>🍽️ Restaurant Finder</h1><p>Pick a town</p></header>'
                                     + f'<div class="towns">{links}</div></body></html>')


if __name__ == "__main__":
    main()
