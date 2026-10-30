"""Find the top rated and most popular restaurants in a town using the Google Places API (New).

Usage:  python3 places_api.py Bhimavaram [Tanuku ...]
Writes results/<Town>_top_rated.xlsx and results/<Town>_most_popular.xlsx.
API key: PLACES_API_KEY env var, or the first line of api_key.txt next to this script.
"""
import math
import os
import sys
import time
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
RESULTS = HERE / "results"
SEARCH_URL = "https://places.googleapis.com/v1/places:searchText"

TOP_N = 50          # restaurants kept per list per town
MIN_RATING = 4.0    # "Most Popular" list only includes places rated at least this
PAD_KM = 1.0        # widen Google's town box by this much on each side (its boxes are tight)
MIN_TILE_KM = 1.5   # stop splitting an area once it is this small
MAX_REQUESTS = int(os.environ.get("MAX_REQUESTS", 400))  # hard cap on billed API calls per run
STATE = "Andhra Pradesh"

# Each query returns at most 60 places, so several phrasings widen coverage; results are deduped by place id
QUERIES = [
    "restaurants", "best restaurants", "biryani", "family restaurant", "veg restaurant",
    "non veg restaurant", "tiffin centre", "cafe", "fast food", "dhaba", "bakery",
    "sea food restaurant", "chinese restaurant", "mess meals", "hotel food",
]

FOOD_TYPES = {
    "restaurant", "food", "cafe", "coffee_shop", "bakery", "meal_takeaway", "meal_delivery",
    "ice_cream_shop", "dessert_shop", "confectionery", "food_court", "juice_shop", "tea_house",
    "bar", "pub", "bar_and_grill", "sandwich_shop", "diner", "deli", "cafeteria", "snack_bar",
}
# Shops Google also tags as "food"/"bakery" but whose main business isn't eating out
NOT_FOOD_PRIMARY = {
    "grocery_store", "supermarket", "hypermarket", "department_store", "shopping_mall",
    "convenience_store", "discount_store", "warehouse_store", "market", "wholesaler", "food_store", "store", "pharmacy", "drugstore",
    "manufacturer",
}

FIELDS = ",".join("places." + f for f in [
    "id", "displayName", "primaryType", "primaryTypeDisplayName", "types", "formattedAddress",
    "nationalPhoneNumber", "rating", "userRatingCount", "priceLevel", "location",
    "googleMapsUri", "businessStatus",
]) + ",nextPageToken"


def load_key():
    key = os.environ.get("PLACES_API_KEY")
    if not key and (HERE / "api_key.txt").exists():
        key = (HERE / "api_key.txt").read_text().strip().splitlines()[0]
    if not key:
        sys.exit("No API key: set PLACES_API_KEY or put it in api_key.txt")
    return key


requests_made = 0


def search(key, body, fields=FIELDS):
    """Run one Text Search, following nextPageToken (max 3 pages / 60 places)."""
    headers = {"X-Goog-Api-Key": key, "X-Goog-FieldMask": fields}
    global requests_made
    places, body = [], dict(body)
    for _ in range(3):
        if requests_made >= MAX_REQUESTS:
            raise RuntimeError(f"Stopped: reached MAX_REQUESTS={MAX_REQUESTS} API calls for this run")
        requests_made += 1
        # 403s can come and go for a while after billing/API changes propagate; 429/5xx are transient too
        for attempt in range(6):
            r = requests.post(SEARCH_URL, json=body, headers=headers, timeout=30)
            if r.status_code not in (403, 429, 500, 503):
                break
            time.sleep(2 * (attempt + 1))
        if r.status_code != 200:
            raise RuntimeError(f"Places API {r.status_code}: {r.text[:500]}")
        data = r.json()
        places += data.get("places", [])
        token = data.get("nextPageToken")
        if not token:
            break
        body["pageToken"] = token
        time.sleep(1)
    return places


def town_area(key, town):
    """Town centre plus its bounding box (Google's viewport for the locality, padded by PAD_KM)."""
    hits = search(key, {"textQuery": f"{town}, {STATE}, India"},
                  "places.location,places.viewport,places.formattedAddress")
    if not hits:
        sys.exit(f"Could not locate {town}")
    loc, vp = hits[0]["location"], hits[0]["viewport"]
    dlat = PAD_KM / 111
    dlng = PAD_KM / (111 * math.cos(math.radians(loc["latitude"])))
    box = (vp["low"]["latitude"] - dlat, vp["low"]["longitude"] - dlng,
           vp["high"]["latitude"] + dlat, vp["high"]["longitude"] + dlng)
    print(f"{town}: centre {loc['latitude']:.4f},{loc['longitude']:.4f}, search box "
          f"{box_km(box)[0]:.1f} x {box_km(box)[1]:.1f} km ({hits[0]['formattedAddress']})")
    return loc["latitude"], loc["longitude"], box


def box_km(box):
    s, w, n, e = box
    return (n - s) * 111, (e - w) * 111 * math.cos(math.radians((n + s) / 2))


def search_area(key, query, box):
    """Search `query` inside box; where Google hits its 60-result cap, split the box in 4 and search each quarter."""
    s, w, n, e = box
    rect = {"rectangle": {"low": {"latitude": s, "longitude": w}, "high": {"latitude": n, "longitude": e}}}
    hits = search(key, {"textQuery": query, "locationRestriction": rect, "languageCode": "en"})
    if len(hits) < 60 or max(box_km(box)) / 2 < MIN_TILE_KM:
        return hits
    mlat, mlng = (s + n) / 2, (w + e) / 2
    for quarter in [(s, w, mlat, mlng), (s, mlng, mlat, e), (mlat, w, n, mlng), (mlat, mlng, n, e)]:
        hits += search_area(key, query, quarter)
    return hits


def distance_m(lat1, lon1, lat2, lon2):
    p1, p2 = math.radians(lat1), math.radians(lat2)
    a = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(lon2 - lon1) / 2) ** 2
    return 6371000 * 2 * math.asin(math.sqrt(a))


def is_food_place(place):
    if place.get("primaryType") in NOT_FOOD_PRIMARY:
        return False
    return any(t in FOOD_TYPES or t.endswith("_restaurant") for t in place.get("types", []))


def collect(key, town):
    lat, lng, box = town_area(key, town)
    found = {}
    for q in QUERIES:
        hits = search_area(key, q, box)
        new = 0
        for p in hits:
            if p["id"] in found or p.get("businessStatus", "OPERATIONAL") != "OPERATIONAL":
                continue
            if not is_food_place(p):
                continue
            loc = p.get("location", {})
            dist = distance_m(lat, lng, loc.get("latitude", lat), loc.get("longitude", lng))
            found[p["id"]] = {
                "Restaurant Name": p.get("displayName", {}).get("text", ""),
                "Category": p.get("primaryTypeDisplayName", {}).get("text", ""),
                "Rating": p.get("rating"),
                "Reviews": p.get("userRatingCount", 0),
                "Price Level": (p.get("priceLevel") or "").replace("PRICE_LEVEL_", "").replace("_", " ").title(),
                "Km from Centre": round(dist / 1000, 1),
                "Address": p.get("formattedAddress", ""),
                "Phone": p.get("nationalPhoneNumber", ""),
                "Google Maps Link": p.get("googleMapsUri", ""),
            }
            new += 1
        print(f"  '{q}': {len(hits)} results, {new} new food places (total {len(found)})")
    return list(found.values())


def add_scores(rows):
    """Bayesian average (IMDb-style weighted rating) so a 5.0 with 3 reviews doesn't beat a 4.6 with 2,000:
    score = v/(v+m) * R + m/(v+m) * C
    R = place rating, v = its review count, C = mean rating in town,
    m = 75th-percentile review count in town (a place needs about this many reviews before its own rating dominates)."""
    rated = [r for r in rows if r["Rating"] is not None and r["Reviews"] > 0]
    C = sum(r["Rating"] for r in rated) / len(rated) if rated else 0
    counts = sorted(r["Reviews"] for r in rated)
    m = max(counts[len(counts) * 3 // 4], 1) if counts else 1
    print(f"IMDb formula: C (town mean rating) = {C:.2f}, m (review threshold) = {m}")
    for r in rows:
        v, R = r["Reviews"], r["Rating"]
        r["Score"] = round(v / (v + m) * R + m / (v + m) * C, 3) if R is not None and v else 0.0


def save(rows, path, title):
    cols = ["Rank", "Restaurant Name", "Category", "Rating", "Reviews", "Score", "Price Level",
            "Km from Centre", "Address", "Phone", "Google Maps Link"]
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(cols)
    for i, r in enumerate(rows, 1):
        ws.append([i] + [r[c] for c in cols[1:]])
        link = ws.cell(row=i + 1, column=len(cols))
        if link.value:
            link.hyperlink = link.value
            link.style = "Hyperlink"
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    widths = {"Restaurant Name": 38, "Category": 22, "Address": 60, "Phone": 16, "Google Maps Link": 45}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 12)
    ws.freeze_panes = "C2"
    wb.save(path)


def run(key, town):
    print(f"\n=== {town} ===")
    rows = collect(key, town)
    if not rows:
        sys.exit(f"No restaurants found for {town}")
    add_scores(rows)
    top_rated = sorted(rows, key=lambda r: r["Score"], reverse=True)[:TOP_N]
    popular = sorted((r for r in rows if (r["Rating"] or 0) >= MIN_RATING),
                     key=lambda r: r["Reviews"], reverse=True)[:TOP_N]
    RESULTS.mkdir(exist_ok=True)
    name = town.replace(" ", "_")
    save(top_rated, RESULTS / f"{name}_top_rated.xlsx", "Top Rated")
    save(popular, RESULTS / f"{name}_most_popular.xlsx", "Most Popular")
    print(f"Google API calls used: {requests_made}")
    print(f"Saved {len(top_rated)} top rated + {len(popular)} most popular to {RESULTS}/")
    return top_rated, popular


if __name__ == "__main__":
    towns = sys.argv[1:] or ["Bhimavaram"]
    key = load_key()
    for t in towns:
        run(key, t)
